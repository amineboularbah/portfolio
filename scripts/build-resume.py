"""Regenerate Word and plain-text downloads from the public Markdown resume.

Requires Python 3.10+ and python-docx 1.2.0. The supplied PDF is never modified.
Only headings, paragraphs, bullets, and inline links are accepted deliberately.
"""
from pathlib import Path
import hashlib
import json
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "src/content/resume.md"
DESTINATION = ROOT / "public/resume"
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def plain(text):
    return LINK.sub(lambda m: m[1] if m[2] == "mailto:" + m[1] else f"{m[1]} ({m[2]})", text)


def add_inline(paragraph, text):
    cursor = 0
    for match in LINK.finditer(text):
        paragraph.add_run(text[cursor:match.start()])
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), paragraph.part.relate_to(match[2], RT.HYPERLINK, is_external=True))
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        properties.extend([color, underline])
        run.append(properties)
        node = OxmlElement("w:t")
        node.text = match[1]
        run.append(node)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        cursor = match.end()
    paragraph.add_run(text[cursor:])


def main():
    source = SOURCE.read_text(encoding="utf-8")
    blocks = source.strip().split("\n\n")
    document = Document()
    document.core_properties.title = "Amine Boularbah"
    document.core_properties.subject = "Senior Mobile Engineer Resume"
    document.core_properties.author = "Amine Boularbah"
    document.core_properties.keywords = "Flutter, Dart, mobile engineering, resume"
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.65)
    section.footer_distance = Inches(0.25)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(10.5)
        style.paragraph_format.line_spacing = 1.05
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.widow_control = True
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.styles["Title"].font.size = Pt(24)
    for name, size in (("Heading 1", 12), ("Heading 2", 10.5)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Amine Boularbah | ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field_run = OxmlElement("w:r")
    field_properties = OxmlElement("w:rPr")
    field_size = OxmlElement("w:sz")
    field_size.set(qn("w:val"), "16")
    field_properties.append(field_size)
    field_run.append(field_properties)
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_run.append(field_text)
    field.append(field_run)
    footer._p.append(field)
    text_blocks = []
    current_section = ""
    for index, block in enumerate(blocks):
        heading = re.fullmatch(r"(#{1,3}) (.+)", block)
        if heading:
            level = len(heading[1])
            title = heading[2]
            if level == 2:
                current_section = title
            paragraph = document.add_paragraph(style={1: "Title", 2: "Heading 1", 3: "Heading 2"}[level])
            if title == "Mobile Engineer at MEGALOGI":
                paragraph.paragraph_format.page_break_before = True
            add_inline(paragraph, title)
            text_blocks.append(plain(title))
        elif block.startswith("- "):
            items = block.splitlines()
            if not all(item.startswith("- ") for item in items):
                raise ValueError("Unsupported mixed Markdown block")
            if current_section == "Skills":
                add_inline(document.add_paragraph(), "; ".join(item[2:] for item in items))
            else:
                for item in items:
                    paragraph = document.add_paragraph(style="List Bullet")
                    add_inline(paragraph, item[2:])
            text_blocks.append("\n".join("- " + plain(item[2:]) for item in items))
        else:
            if re.search(r"^\s*[#>*`|]|[!*`]", block, re.M):
                raise ValueError("Unsupported Markdown syntax; update the exporter before continuing")
            paragraph = document.add_paragraph()
            add_inline(paragraph, " ".join(block.splitlines()))
            if index == 1:
                paragraph.runs[0].font.size = Pt(12)
                paragraph.runs[0].bold = True
            if index in (2, 3):
                paragraph.style = document.styles["Normal"]
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            if index > 0 and blocks[index - 1].startswith("### "):
                paragraph.paragraph_format.keep_with_next = current_section in ("Employment History", "Products")
            text_blocks.append(plain(" ".join(block.splitlines())))
    DESTINATION.mkdir(exist_ok=True)
    document.save(DESTINATION / "amine-boularbah-resume.docx")
    (DESTINATION / "amine-boularbah-resume.txt").write_text("\n\n".join(text_blocks) + "\n", encoding="utf-8")
    paths = [SOURCE, *(DESTINATION / f"amine-boularbah-resume.{extension}" for extension in ("pdf", "docx", "txt"))]
    manifest = {str(path.relative_to(ROOT)): hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
    (ROOT / "scripts/resume-manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print("Updated Word, text, and checksum manifest. Original PDF preserved.")


if __name__ == "__main__":
    main()
